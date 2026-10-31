#!/usr/bin/env python3.13
"""Build the TenantPro project report.

Front matter follows the four front pages of the supplied M.Com report; the body
follows the chapter structure of the supplied project-report format.

Two passes: pass 1 writes the document with blank page numbers, renders it to PDF,
and reads the real page of every heading; pass 2 writes it again with those numbers
in the contents, list of tables and list of figures.
"""
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
CONTENT = os.path.join(HERE, "content.txt")

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(12)
LINE = 1.5

TITLE = "TENANTPRO — SMART TENANT AND RENT MANAGEMENT SYSTEM"
TITLE_PLAIN = "TenantPro — Smart Tenant and Rent Management System"
STUDENTS = [
    ("Anush Kulal M", "24MCAR0133"),
    ("Akshita Bansal", "24MCAR0131"),
    ("Chandan Mathur", "24MCAR0230"),
]
GUIDE = "Dr. Kamalraj R"
GUIDE_DESIG = "Department of Computer Science and IT"
DIRECTOR = "Dr. SAGAR GULATI"
SCHOOL = "SCHOOL OF COMPUTER SCIENCE AND IT"
UNIV = "JAIN (Deemed-to-be University)"
PLACE = "Bengaluru"
YEAR = "2025 - 2026"


# ----------------------------------------------------------------- docx helpers
def set_base_style(doc):
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = BODY_SIZE
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = st.paragraph_format
    pf.line_spacing = LINE
    pf.space_after = Pt(6)


def para(doc, text="", *, align=None, bold=False, italic=False, size=None,
         caps=False, space_before=None, space_after=None, line=None,
         indent_first=None, keep_with_next=False):
    p = doc.add_paragraph()
    run = p.add_run(text.upper() if caps else text)
    run.bold = bold
    run.italic = italic
    run.font.name = BODY_FONT
    run.font.size = size or BODY_SIZE
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = line if line is not None else LINE
    if space_before is not None:
        pf.space_before = space_before
    if space_after is not None:
        pf.space_after = space_after
    if indent_first is not None:
        pf.first_line_indent = indent_first
    pf.keep_with_next = keep_with_next
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def blank_lines(doc, n=1, size=Pt(12)):
    for _ in range(n):
        para(doc, "", size=size, line=1.0, space_after=Pt(0))


def add_page_number_footer(section, fmt="decimal", start=None):
    """Centered PAGE field in the footer."""
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = ""
    run = p.add_run()
    run.font.name = BODY_FONT
    run.font.size = Pt(11)
    for el, attrs, text in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, " PAGE "),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        node = OxmlElement(el)
        for k, v in attrs.items():
            node.set(qn(k) if ":" in k else k, v)
        if text:
            node.text = text
        run._r.append(node)
    sect_pr = section._sectPr
    pg = sect_pr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sect_pr.append(pg)
    pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))


def page_border(section, color="1F3864", sz=14, space=22):
    """A frame on every page of this section, measured from the page edge."""
    sect_pr = section._sectPr
    old = sect_pr.find(qn("w:pgBorders"))
    if old is not None:
        sect_pr.remove(old)
    pb = OxmlElement("w:pgBorders")
    pb.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), str(space))
        el.set(qn("w:color"), color)
        pb.append(el)
    sect_pr.append(pb)


def different_first_page(section):
    """No page number on the title page."""
    sect_pr = section._sectPr
    el = sect_pr.find(qn("w:titlePg"))
    if el is None:
        el = OxmlElement("w:titlePg")
        sect_pr.append(el)


def same_first_page(section):
    """python-docx copies the previous sectPr, so w:titlePg is inherited. The body
    section needs its own first page numbered."""
    sect_pr = section._sectPr
    el = sect_pr.find(qn("w:titlePg"))
    if el is not None:
        sect_pr.remove(el)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    tcPr.append(el)


SMALL_WORDS = {"and", "or", "of", "the", "a", "an", "in", "on", "for", "to",
               "by", "with", "at", "from", "not"}


def nice_title(text):
    words = text.split()
    out = []
    for i, w in enumerate(words):
        lw = w.lower()
        if i and lw in SMALL_WORDS:
            out.append(lw)
        elif w.isupper() and len(w) > 3:
            out.append(w.capitalize())
        else:
            out.append(w if not w.isupper() else w.capitalize())
    return " ".join(out)


def set_widths(t, widths):
    """Fix the column widths. Word reads w:tblGrid as well as the cell widths,
    so both have to be written or the table falls back to equal columns."""
    t.autofit = False
    tbl = t._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(sum(w.twips for w in widths)))
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for col, w in zip(grid.findall(qn("w:gridCol")), widths):
            col.set(qn("w:w"), str(w.twips))
    for row in t.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w


def add_table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = BODY_FONT
            run.font.size = Pt(10.5)
            run.bold = (r == 0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            if r == 0:
                shade(cell, "E8E8E8")
        if r == 0:
            # repeat the header when the table runs over a page
            trPr = t.rows[0]._tr.get_or_add_trPr()
            hdr = OxmlElement("w:tblHeader")
            trPr.append(hdr)
        else:
            trPr = t.rows[r]._tr.get_or_add_trPr()
            keep = OxmlElement("w:cantSplit")
            trPr.append(keep)
    return t


def auto_widths(rows, total=Inches(6.05)):
    """Share the width out between the columns.

    Each column first gets enough room for its longest unbreakable word, since a
    column narrower than that breaks AUTO_INCREMENT across two lines. Whatever is
    left over is then shared in proportion to how much text each column carries.
    """
    ncol = len(rows[0])
    char = 0.084   # inches per character at 10.5 pt; capitals run wider
    pad = 0.16     # cell padding either side

    need, mean = [], []
    for c in range(ncol):
        longest_word = max((len(w) for r in rows for w in r[c].split()), default=4)
        need.append(min(longest_word * char + pad, 1.9))
        mean.append(max(sum(len(r[c]) for r in rows) / len(rows), 4.0))

    spare = total.twips - sum(Inches(n).twips for n in need)
    if spare <= 0:                       # nothing to share out; scale back to fit
        k = total.twips / sum(Inches(n).twips for n in need)
        return [Inches(n * k) for n in need]
    scale = spare / sum(mean)
    return [Inches(n + m * scale / 1440) for n, m in zip(need, mean)]


# ------------------------------------------------------------------ front pages
LOGO = os.path.join(ASSETS, "logo_cs.png")


def put_logo(doc, width=Inches(2.6), after=Pt(10)):
    if not os.path.exists(LOGO):
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = after
    p.add_run().add_picture(LOGO, width=width)


def title_page(doc):
    blank_lines(doc, 2)
    para(doc, TITLE, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=Pt(18), line=1.3)
    blank_lines(doc, 1)
    put_logo(doc, width=Inches(3.4))
    blank_lines(doc, 1)
    for line in (
        "A project report submitted in partial fulfilment of the requirements of",
        "Master of Computer Applications Post Graduate Course of",
        UNIV,
    ):
        para(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(13), line=1.3, space_after=Pt(2))
    blank_lines(doc, 1)
    para(doc, "Submitted by", align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(13), space_after=Pt(4))
    for name, usn in STUDENTS:
        para(doc, f"{name}        {usn}", align=WD_ALIGN_PARAGRAPH.CENTER,
             bold=True, size=Pt(13), line=1.2, space_after=Pt(2))
    blank_lines(doc, 1)
    para(doc, "Under the guidance of", align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(13), space_after=Pt(4))
    para(doc, GUIDE, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=Pt(13), space_after=Pt(2))
    for line in (GUIDE_DESIG, UNIV, PLACE):
        para(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(13), line=1.2, space_after=Pt(2))
    blank_lines(doc, 1)
    para(doc, "Date -", align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(13))


def certificate_university(doc):
    page_break(doc)
    put_logo(doc, width=Inches(2.4), after=Pt(18))
    para(doc, "CERTIFICATE BY THE UNIVERSITY", align=WD_ALIGN_PARAGRAPH.CENTER,
         bold=True, size=Pt(15))
    blank_lines(doc, 2)
    names = ", ".join(f"{n} (Register No. {u})" for n, u in STUDENTS)
    para(doc,
         f"Certified that the project report titled “{TITLE_PLAIN}” is based on an "
         f"original project study conducted by {names} of the Master of Computer Applications "
         f"programme, under the guidance of {GUIDE}. They have attended the required guidance "
         f"sessions held. The project report has not formed a basis for the award of any "
         f"Degree or Diploma of any university or any other institution.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=2.0)
    blank_lines(doc, 3)
    para(doc, f"PLACE: {PLACE}", space_after=Pt(2))
    para(doc, "DATE:")
    blank_lines(doc, 3)
    for line in (DIRECTOR, "DIRECTOR", SCHOOL, UNIV):
        para(doc, line, align=WD_ALIGN_PARAGRAPH.RIGHT, bold=(line == DIRECTOR),
             line=1.2, space_after=Pt(2))


def certificate_guide(doc):
    page_break(doc)
    put_logo(doc, width=Inches(2.4), after=Pt(18))
    para(doc, "CERTIFICATE BY THE GUIDE", align=WD_ALIGN_PARAGRAPH.CENTER,
         bold=True, size=Pt(15))
    blank_lines(doc, 2)
    names = ", ".join(f"{n} (Register No. {u})" for n, u in STUDENTS)
    para(doc,
         f"Certified that the project report titled “{TITLE_PLAIN}” is based on an "
         f"original project study conducted by {names} under my guidance. This project has not "
         f"formed a basis for the award of any Degree or Diploma of any University or any other "
         f"Institution.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=2.0)
    blank_lines(doc, 3)
    para(doc, f"PLACE: {PLACE}", space_after=Pt(2))
    para(doc, "DATE:")
    blank_lines(doc, 4)
    para(doc, "Signature of the Guide", align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=Pt(2))
    para(doc, GUIDE, align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True, space_after=Pt(2))
    para(doc, GUIDE_DESIG, align=WD_ALIGN_PARAGRAPH.RIGHT, line=1.2)


def declaration(doc):
    page_break(doc)
    put_logo(doc, width=Inches(2.4), after=Pt(18))
    para(doc, "DECLARATION BY THE STUDENTS", align=WD_ALIGN_PARAGRAPH.CENTER,
         bold=True, size=Pt(15))
    blank_lines(doc, 2)
    para(doc,
         f"This is to declare that the project titled “{TITLE_PLAIN}” is carried out "
         f"independently by us under the guidance of {GUIDE}, {GUIDE_DESIG}, {UNIV}, {PLACE}. "
         f"This work is an original one and has not been submitted earlier to any University or "
         f"any other Institution for the fulfilment of the requirement of a course of study or "
         f"any other credential.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=2.0)
    blank_lines(doc, 3)
    para(doc, f"PLACE: {PLACE}", space_after=Pt(2))
    para(doc, "DATE:")
    blank_lines(doc, 3)
    for name, usn in STUDENTS:
        para(doc, f"NAME: {name}", align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True,
             line=1.2, space_after=Pt(0))
        para(doc, f"USN: {usn}", align=WD_ALIGN_PARAGRAPH.RIGHT, line=1.2, space_after=Pt(8))


def acknowledgement(doc):
    page_break(doc)
    put_logo(doc, width=Inches(2.4), after=Pt(18))
    para(doc, "ACKNOWLEDGEMENT", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=Pt(15))
    blank_lines(doc, 1)
    body = [
        f"We are deeply indebted to Dr. CHENRAJ ROYCHAND, President, {UNIV} Trust, {PLACE}, for "
        f"having admitted us to undergo the Master of Computer Applications Post Graduate course "
        f"during the academic term {YEAR}.",
        f"We are pleased to thank {DIRECTOR.title()}, Director, School of Computer Science and IT, "
        f"{UNIV}, {PLACE}, for the facilities and the academic environment extended to us during "
        f"the course of this work.",
        f"We take this opportunity to express our gratitude to Dr. Ananta Charan Ojha, Deputy "
        f"Director, Dr. K Suneetha, Head of the Department of Computer Science and IT, and "
        f"Dr. Murugan R, Programme Head of the MCA Programme, for their valuable guidance and "
        f"support for the successful completion of the project work.",
        f"We express our profound thanks to our project coordinator Dr. I Brem Navas, Professor, "
        f"Department of Computer Science and IT, and above all to our guide {GUIDE}, for guidance "
        f"at the decisive points of this work and for being available whenever it mattered.",
        "We are very thankful to our families and friends for their constant and unconditional "
        "encouragement and support.",
    ]
    for t in body:
        para(doc, t, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    blank_lines(doc, 2)
    for name, usn in STUDENTS:
        para(doc, f"NAME: {name}", align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True,
             line=1.2, space_after=Pt(0))
        para(doc, f"USN: {usn}", align=WD_ALIGN_PARAGRAPH.RIGHT, line=1.2, space_after=Pt(8))


# ------------------------------------------------------------------- content io
def parse_content(path):
    items = []
    for raw in open(path, encoding="utf-8"):
        line = raw.rstrip("\n")
        if not line.strip():
            continue
        if line.startswith("|"):
            items.append(("ROW", [c.strip() for c in line[1:].split("|")]))
            continue
        tag, _, rest = line.partition(" ")
        items.append((tag[1:], rest))
    return items


# ---------------------------------------------------------------------- builder
def build(out_path, pagemap=None):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Inches(11.69)
    sec.page_width = Inches(8.27)
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.25)
    sec.right_margin = Inches(1.0)
    set_base_style(doc)
    page_border(sec)
    different_first_page(sec)
    add_page_number_footer(sec, fmt="lowerRoman", start=1)

    title_page(doc)
    certificate_university(doc)
    certificate_guide(doc)
    declaration(doc)
    acknowledgement(doc)

    items = parse_content(CONTENT)

    # collect headings, tables and figures for the three lists
    toc, tables, figures = [], [], []
    for tag, val in items:
        if tag == "CH":
            toc.append(("chapter", val))
        elif tag == "S":
            toc.append(("section", val))
        elif tag == "ABSTRACT":
            pass
        elif tag == "TBL":
            tables.append(val)
        elif tag == "FIG":
            figures.append(val.split("|", 1)[1])
    toc = [("chapter", "ABSTRACT")] + toc

    def pg(key):
        if not pagemap:
            return ""
        return str(pagemap.get(key, ""))

    # ---- table of contents
    page_break(doc)
    para(doc, "TABLE OF CONTENTS", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=Pt(15))
    blank_lines(doc, 1)
    rows = [["Chapter", "Title", "Page No."]]
    for kind, text in toc:
        if kind == "chapter":
            m = re.match(r"CHAPTER (\d+) — (.+)", text)
            if m:
                rows.append([m.group(1), nice_title(m.group(2)), pg(text)])
            else:
                rows.append(["", nice_title(text), pg(text)])
        else:
            num, _, rest = text.partition(" ")
            rows.append(["", f"        {num}  {rest}", pg(text)])
    t = add_table(doc, rows)
    set_widths(t, [Inches(0.85), Inches(4.35), Inches(0.85)])

    # ---- list of tables
    page_break(doc)
    para(doc, "LIST OF TABLES", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=Pt(15))
    blank_lines(doc, 1)
    rows = [["Table No.", "Title", "Page No."]]
    for cap in tables:
        num, _, name = cap.partition(" — ")
        rows.append([num.replace("Table ", ""), name, pg(cap)])
    t = add_table(doc, rows)
    set_widths(t, [Inches(0.85), Inches(4.35), Inches(0.85)])

    # ---- list of figures
    page_break(doc)
    para(doc, "LIST OF FIGURES", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=Pt(15))
    blank_lines(doc, 1)
    rows = [["Figure No.", "Title", "Page No."]]
    for cap in figures:
        num, _, name = cap.partition(" — ")
        rows.append([num.replace("Fig ", ""), name, pg(cap)])
    t = add_table(doc, rows)
    set_widths(t, [Inches(0.85), Inches(4.35), Inches(0.85)])

    # ---- body in its own section, page numbers restart at 1
    body_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    body_sec.page_height = Inches(11.69)
    body_sec.page_width = Inches(8.27)
    body_sec.top_margin = Inches(1.0)
    body_sec.bottom_margin = Inches(1.0)
    body_sec.left_margin = Inches(1.25)
    body_sec.right_margin = Inches(1.0)
    page_border(body_sec)
    same_first_page(body_sec)
    add_page_number_footer(body_sec, start=1)

    pending_rows = []

    def flush_table(caption):
        nonlocal pending_rows
        if not pending_rows:
            return
        tb = add_table(doc, pending_rows)
        set_widths(tb, auto_widths(pending_rows))
        para(doc, "", size=Pt(6), line=1.0, space_after=Pt(0))
        pending_rows = []

    current_caption = None
    first_block = True
    for idx, (tag, val) in enumerate(items):
        if tag == "ROW":
            pending_rows.append(val)
            continue
        if pending_rows:
            flush_table(current_caption)

        if tag == "ABSTRACT":
            if not first_block:
                page_break(doc)
            para(doc, "ABSTRACT", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=Pt(16),
                 space_after=Pt(14), keep_with_next=True)
            first_block = False
        elif tag == "CH":
            if not first_block:
                page_break(doc)
            para(doc, val, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=Pt(16),
                 space_after=Pt(14), keep_with_next=True)
            first_block = False
        elif tag == "S":
            para(doc, val, bold=True, size=Pt(13.5), space_before=Pt(12),
                 space_after=Pt(6), keep_with_next=True)
        elif tag == "SS":
            para(doc, val, bold=True, size=Pt(12), space_before=Pt(10),
                 space_after=Pt(4), keep_with_next=True)
        elif tag == "P":
            para(doc, val, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif tag == "N":
            p = doc.add_paragraph(val, style="List Number")
            for r in p.runs:
                r.font.name = BODY_FONT
                r.font.size = BODY_SIZE
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = LINE
        elif tag == "B":
            p = doc.add_paragraph(val, style="List Bullet")
            for r in p.runs:
                r.font.name = BODY_FONT
                r.font.size = BODY_SIZE
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = LINE
        elif tag == "R":
            p = para(doc, val, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.15,
                     space_after=Pt(8))
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.first_line_indent = Inches(-0.4)
        elif tag == "TBL":
            current_caption = val
            para(doc, val, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=Pt(11),
                 space_before=Pt(10), space_after=Pt(4), line=1.15, keep_with_next=True)
        elif tag == "FIG":
            fname, caption = val.split("|", 1)
            path = os.path.join(ASSETS, fname)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            import struct
            with open(path, "rb") as fh:
                head = fh.read(33)
            w, h = struct.unpack(">II", head[16:24])
            # drawn figures may run tall and get a page to themselves; phone
            # screenshots are capped shorter so they sit beside their text
            max_w = 5.85
            max_h = 8.5 if os.path.basename(fname).startswith("fig_") else 5.6
            scale = min(max_w / w, max_h / h)
            p.add_run().add_picture(path, width=Inches(w * scale))
            para(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=Pt(11),
                 line=1.15, space_after=Pt(12))

    if pending_rows:
        flush_table(current_caption)

    doc.save(out_path)
    return toc, tables, figures


def render_pdf(docx_path, outdir):
    subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, docx_path],
        check=True, capture_output=True, timeout=600,
    )
    return os.path.join(outdir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")


def page_numbers(pdf_path, toc, tables, figures):
    """Map each heading/caption to the printed page number of the body section."""
    txt = subprocess.run(["pdftotext", "-layout", pdf_path, "-"],
                         check=True, capture_output=True, text=True).stdout
    pages = txt.split("\f")
    # the body starts on the first page that carries the ABSTRACT heading
    body_start = next(i for i, pg in enumerate(pages) if "ABSTRACT" in pg)

    def norm(s):
        return re.sub(r"\s+", " ", s).strip().lower()

    mapping = {}
    targets = []
    for kind, text in toc:
        targets.append((text, norm(text)))
    for cap in tables + figures:
        # captions wrap across lines in the rendered page; a prefix is enough to
        # locate them and survives the wrap
        targets.append((cap, norm(cap)[:40]))

    normpages = [norm(p) for p in pages]
    tops = []
    for p in pages:
        head = [l.strip() for l in p.split("\n") if l.strip()][:3]
        tops.append(norm(" ".join(head)))

    chapter_keys = {t for k, t in toc if k == "chapter"}
    for key, needle in targets:
        if not needle:
            continue
        pool = tops if key in chapter_keys else normpages
        for i in range(body_start, len(pages)):
            if needle in pool[i]:
                mapping[key] = i - body_start + 1
                break
    return mapping


def main():
    out_dir = sys.argv[1] if len(sys.argv) > 1 else HERE
    tmp = os.path.join(HERE, "pass1.docx")
    toc, tables, figures = build(tmp)
    pdf = render_pdf(tmp, HERE)
    mapping = page_numbers(pdf, toc, tables, figures)
    missing = [k for k, _ in [(t, None) for _, t in toc]] if False else None
    final = os.path.join(out_dir, "TenantPro Project Report.docx")
    build(final, pagemap=mapping)
    print("pass 1:", tmp)
    print("mapped", len(mapping), "of", len(toc) + len(tables) + len(figures), "entries")
    unmapped = [t for _, t in toc if t not in mapping]
    if unmapped:
        print("UNMAPPED:", unmapped)
    print("final:", final)


if __name__ == "__main__":
    main()
